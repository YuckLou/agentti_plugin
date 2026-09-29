#!/usr/bin/env python3
"""Gera o relatório Agentti em Word (.docx) e PDF a partir de um JSON de conteúdo.

    python gerar_relatorio.py conteudo.json [--saida PASTA] [--previa]

O formato do JSON está em ../references/formato.md. O visual (cores, cabeçalho com logo, rodapé, tabelas,
gráficos) é fixo aqui, para sair igual em qualquer lugar; o Claude só escreve o conteúdo.

Saída: uma linha JSON com os caminhos gerados, ou a lista de erros do conteúdo (código 2) ou dos pacotes que
faltam (código 3).
"""
import argparse
import json
import os
import re
import sys
import tempfile
from datetime import date

AQUI = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(AQUI, "..", "assets", "logo.png")

# Tema claro da marca (landing-agentti, [data-theme="light"]).
LARANJA, GRAFITE, TEXTO, SECUNDARIO, FUNDO, CINZA = "D84315", "263238", "0F172A", "475569", "F1F5F9", "94A3B8"
LINHA = "E2E8F0"
SERIES = ["#" + LARANJA, "#" + GRAFITE, "#" + CINZA]
# Arial em vez de Chakra Petch/Inter: abre igual em qualquer máquina e no Google Docs.
FONTE = "Arial"

RODAPE = "Dados: base oficial da Receita Federal · análise Agentti · agentti.ia.br"
SOBRE = ("**Sobre os dados.** Os números vêm da contagem completa dos estabelecimentos ativos na base oficial da "
         "Receita Federal, consultada pelo Agentti em {data}. A faixa A/B/C é uma pré-qualificação feita só com o "
         "cadastro (porte, tempo ativa, canal de contato e nome fantasia) e não indica se a empresa compra o "
         "produto.{contatos} Capital social é o declarado na abertura e não representa faturamento. "
         "Mais em agentti.ia.br.")
SOBRE_CONTATOS = (" Contatos marcados como \"confirmado\" foram achados numa página que traz o endereço ou o "
                  "telefone da Receita; \"provável\" indica sinais sem essa prova.")

BLOCOS = ("texto", "lista", "tabela", "grafico", "grafico_empilhado")


# ── Validação ────────────────────────────────────────────────────────────────

def validar(c):
    erros = []
    if not isinstance(c, dict):
        return ["o conteúdo tem de ser um objeto JSON"]
    for k in ("titulo", "secoes"):
        if not c.get(k):
            erros.append(f"falta '{k}'")
    for i, d in enumerate(c.get("destaques") or []):
        if not isinstance(d, dict) or "valor" not in d or "rotulo" not in d:
            erros.append(f"destaques[{i}]: precisa de 'valor' e 'rotulo'")
    if len(c.get("destaques") or []) > 4:
        erros.append("destaques: no máximo 4")
    for i, s in enumerate(c.get("secoes") or []):
        onde = f"secoes[{i}]"
        if not s.get("titulo"):
            erros.append(f"{onde}: falta 'titulo'")
        for j, b in enumerate(s.get("blocos") or []):
            ob = f"{onde}.blocos[{j}]"
            chaves = [k for k in b if k in BLOCOS] if isinstance(b, dict) else []
            if len(chaves) != 1:
                erros.append(f"{ob}: use exatamente uma destas chaves: {', '.join(BLOCOS)}")
                continue
            tipo, v = chaves[0], b[chaves[0]]
            if tipo == "tabela":
                cols, linhas = v.get("colunas") or [], v.get("linhas") or []
                if not cols or not linhas:
                    erros.append(f"{ob}: tabela precisa de 'colunas' e 'linhas'")
                for n, ln in enumerate(linhas):
                    if len(ln) != len(cols):
                        erros.append(f"{ob}: a linha {n} tem {len(ln)} células e há {len(cols)} colunas")
            elif tipo == "grafico":
                rot, val = v.get("rotulos") or [], v.get("valores") or []
                if not rot or len(rot) != len(val):
                    erros.append(f"{ob}: 'rotulos' e 'valores' precisam ter o mesmo tamanho (e não vazios)")
                if any(not isinstance(x, (int, float)) for x in val):
                    erros.append(f"{ob}: 'valores' só com números (sem texto, sem '%')")
            elif tipo == "grafico_empilhado":
                cats, series = v.get("categorias") or [], v.get("series") or {}
                if not cats or not series:
                    erros.append(f"{ob}: precisa de 'categorias' e 'series'")
                for nome, vals in series.items():
                    if len(vals) != len(cats) or any(not isinstance(x, (int, float)) for x in vals):
                        erros.append(f"{ob}: a série '{nome}' precisa de {len(cats)} números")
            elif tipo == "lista" and not isinstance(v, list):
                erros.append(f"{ob}: 'lista' é uma lista de textos")
    return erros


# ── Números e gráficos ───────────────────────────────────────────────────────

def num_br(x, casas=0):
    s = f"{x:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def _pct(v, total):
    return f"{num_br(100.0 * v / total, 1)}%" if total else ""


def grafico_barras(g, caminho):
    import matplotlib
    matplotlib.use("Agg")
    # Arial; no Linux, a Liberation Sans (mesmas medidas). Sem as duas, a DejaVu do matplotlib.
    matplotlib.rcParams["font.family"] = "sans-serif"
    matplotlib.rcParams["font.sans-serif"] = ["Arial", "Liberation Sans", "DejaVu Sans"]
    import matplotlib.pyplot as plt
    rot, val = list(g["rotulos"]), list(g["valores"])
    total = g.get("total") or sum(val)
    fig, ax = plt.subplots(figsize=(7.2, 0.42 * len(rot) + 0.6), dpi=160)
    y = range(len(rot))[::-1]
    ax.barh(list(y), val, color="#" + LARANJA, height=0.62)
    ax.set_yticks(list(y), rot, fontsize=9, color="#" + TEXTO)
    topo = max(val) if val else 1
    for yi, v in zip(y, val):
        ax.text(v + topo * 0.01, yi, f"{num_br(v)}  ({_pct(v, total)})", va="center", fontsize=8.5, color="#" + SECUNDARIO)
    ax.set_xlim(0, topo * 1.22)
    ax.xaxis.set_visible(False)
    for lado in ("top", "right", "bottom"):
        ax.spines[lado].set_visible(False)
    ax.spines["left"].set_color("#" + LINHA)
    ax.tick_params(axis="y", length=0)
    fig.tight_layout()
    fig.savefig(caminho, facecolor="white")
    plt.close(fig)


def grafico_empilhado(g, caminho):
    import matplotlib
    matplotlib.use("Agg")
    # Arial; no Linux, a Liberation Sans (mesmas medidas). Sem as duas, a DejaVu do matplotlib.
    matplotlib.rcParams["font.family"] = "sans-serif"
    matplotlib.rcParams["font.sans-serif"] = ["Arial", "Liberation Sans", "DejaVu Sans"]
    import matplotlib.pyplot as plt
    cats, series = list(g["categorias"]), g["series"]
    fig, ax = plt.subplots(figsize=(7.2, 0.5 * len(cats) + 0.9), dpi=160)
    y = list(range(len(cats)))[::-1]
    base = [0.0] * len(cats)
    tot = [sum(vals[i] for vals in series.values()) for i in range(len(cats))]
    # Cada categoria em 100%: em números absolutos as categorias pequenas (ex.: EPP ao lado do MEI) somem.
    for k, (nome, vals) in enumerate(series.items()):
        cor = SERIES[k % len(SERIES)]
        pct = [100.0 * v / t if t else 0 for v, t in zip(vals, tot)]
        ax.barh(y, pct, left=base, color=cor, height=0.62, label=nome)
        for i, (v, p) in enumerate(zip(vals, pct)):
            if p >= 9:
                ax.text(base[i] + p / 2, y[i], num_br(v), ha="center", va="center", fontsize=8,
                        color="white" if k < 2 else "#" + TEXTO)
        base = [b + p for b, p in zip(base, pct)]
    for i, t in enumerate(tot):
        ax.text(101, y[i], num_br(t), va="center", fontsize=8, color="#" + SECUNDARIO)
    ax.set_xlim(0, 110)
    ax.set_yticks(y, cats, fontsize=9, color="#" + TEXTO)
    ax.xaxis.set_visible(False)
    for lado in ("top", "right", "bottom", "left"):
        ax.spines[lado].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.legend(ncol=len(series), frameon=False, fontsize=8.5, loc="lower center", bbox_to_anchor=(0.5, 1.0))
    fig.tight_layout()
    fig.savefig(caminho, facecolor="white")
    plt.close(fig)


def _numerico(v):
    return bool(re.fullmatch(r"[\d.,%R$\s|-]+", str(v)))


def _colunas_numericas(t):
    """Colunas (menos a primeira) em que todo valor é número: alinham à direita, cabeçalho junto."""
    return {j for j in range(1, len(t["colunas"])) if t["linhas"] and all(_numerico(ln[j]) for ln in t["linhas"])}


def tabela_do_grafico(g):
    total = g.get("total") or sum(g["valores"])
    linhas = [[r, num_br(v), _pct(v, total)] for r, v in zip(g["rotulos"], g["valores"])]
    return {"colunas": [g.get("rotulo_categoria", "Categoria"), g.get("unidade", "Empresas"), "% do total"], "linhas": linhas}


def tabela_do_empilhado(g):
    nomes = list(g["series"])
    linhas = []
    for i, cat in enumerate(g["categorias"]):
        vals = [g["series"][n][i] for n in nomes]
        linhas.append([cat] + [num_br(v) for v in vals] + [num_br(sum(vals))])
    return {"colunas": [g.get("rotulo_categoria", "Categoria")] + nomes + ["Total"], "linhas": linhas}


# ── Texto com **negrito** ────────────────────────────────────────────────────

def partes(texto):
    """[(trecho, negrito)] a partir de '**...**'."""
    out = []
    for i, p in enumerate(re.split(r"\*\*", str(texto))):
        if p:
            out.append((p, i % 2 == 1))
    return out


# ── Word ─────────────────────────────────────────────────────────────────────

def gerar_docx(c, graficos, caminho):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    cor = lambda h: RGBColor.from_string(h)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2)
    largura = Cm(17)

    def estilo(nome, tam, cor_hex, negrito=False, italico=False, antes=0, depois=6):
        st = doc.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        st.font.name, st.font.size, st.font.bold, st.font.italic = FONTE, Pt(tam), negrito, italico
        st.font.color.rgb = cor(cor_hex)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(antes), Pt(depois)
        return st

    doc.styles["Normal"].font.name = FONTE
    estilo("Título Agentti", 20, LARANJA, negrito=True, depois=2)
    estilo("Subtítulo Agentti", 9, SECUNDARIO, depois=12)
    estilo("Seção Agentti", 14, LARANJA, negrito=True, antes=14, depois=6).paragraph_format.keep_with_next = True
    estilo("Texto Agentti", 10.5, TEXTO, depois=6)
    estilo("Legenda Agentti", 8.5, SECUNDARIO, italico=True, antes=2, depois=10)
    estilo("Célula Agentti", 9, TEXTO, depois=0)
    estilo("Destaque Agentti", 18, LARANJA, negrito=True, depois=0).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    estilo("Rodapé Agentti", 8, SECUNDARIO, depois=0)

    def sombra(cell, fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
        tcpr.append(shd)

    def bordas(tabela):
        tblpr = tabela._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{lado}")
            e.set(qn("w:val"), "single"), e.set(qn("w:sz"), "4"), e.set(qn("w:color"), LINHA)
            b.append(e)
        tblpr.append(b)

    def paragrafo(texto, estilo_nome, alvo=None):
        p = (alvo or doc).add_paragraph(style=estilo_nome)
        for trecho, neg in partes(texto):
            p.add_run(trecho).bold = neg or None
        return p

    def tabela(t):
        cols, linhas = t["colunas"], t["linhas"]
        num = _colunas_numericas(t)
        tb = doc.add_table(rows=1 + len(linhas), cols=len(cols))
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordas(tb)
        for j, nome in enumerate(cols):
            cel = tb.rows[0].cells[j]
            cel.text = ""
            r = cel.paragraphs[0].add_run(str(nome))
            cel.paragraphs[0].style = "Célula Agentti"
            if j in num:
                cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r.bold, r.font.color.rgb = True, cor("FFFFFF")
            sombra(cel, GRAFITE)
        for i, ln in enumerate(linhas, start=1):
            for j, v in enumerate(ln):
                cel = tb.rows[i].cells[j]
                cel.text = ""
                p = cel.paragraphs[0]
                p.style = "Célula Agentti"
                for trecho, neg in partes(v):
                    p.add_run(trecho).bold = neg or None
                if j in num:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if i % 2 == 0:
                    sombra(cel, FUNDO)
        rpr = tb.rows[0]._tr.get_or_add_trPr()
        cab = OxmlElement("w:tblHeader")
        cab.set(qn("w:val"), "true")
        rpr.append(cab)
        if t.get("legenda"):
            paragrafo(t["legenda"], "Legenda Agentti")
        else:
            doc.add_paragraph(style="Legenda Agentti")

    # Cabeçalho: logo pequeno à esquerda, título curto à direita. Rodapé: origem, data e página.
    cab = sec.header.paragraphs[0]
    cab.paragraph_format.tab_stops.add_tab_stop(largura, WD_TAB_ALIGNMENT.RIGHT)
    if os.path.exists(LOGO):
        cab.add_run().add_picture(LOGO, width=Cm(2.2))
    r = cab.add_run("\t" + (c.get("titulo_curto") or c["titulo"]))
    r.font.size, r.font.color.rgb, r.font.name = Pt(8), cor(SECUNDARIO), FONTE
    rod = sec.footer.paragraphs[0]
    rod.style = "Rodapé Agentti"
    rod.paragraph_format.tab_stops.add_tab_stop(largura, WD_TAB_ALIGNMENT.RIGHT)
    rod.add_run(f"{RODAPE} · gerado em {c['data']}\tPágina ")
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    rr, tt = OxmlElement("w:r"), OxmlElement("w:t")
    tt.text = "1"
    rr.append(tt), campo.append(rr)
    rod._p.append(campo)

    paragrafo(c["titulo"], "Título Agentti")
    if c.get("subtitulo"):
        paragrafo(c["subtitulo"], "Subtítulo Agentti")

    if c.get("destaques"):
        ds = c["destaques"]
        tb = doc.add_table(rows=2, cols=len(ds))
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, d in enumerate(ds):
            for i, (txt, st) in enumerate(((d["valor"], "Destaque Agentti"), (d["rotulo"], "Célula Agentti"))):
                cel = tb.rows[i].cells[j]
                cel.text = ""
                cel.paragraphs[0].style = st
                cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cel.paragraphs[0].add_run(str(txt))
                sombra(cel, FUNDO)
        doc.add_paragraph(style="Legenda Agentti")

    for n, s in enumerate(c["secoes"], start=1):
        titulo = s["titulo"] if c.get("numerar") is False else f"{n}. {s['titulo']}"
        paragrafo(titulo, "Seção Agentti")
        for b in s.get("blocos") or []:
            if "texto" in b:
                paragrafo(b["texto"], "Texto Agentti")
            elif "lista" in b:
                for item in b["lista"]:
                    p = paragrafo("• " + str(item), "Texto Agentti")
                    p.paragraph_format.left_indent = Cm(0.5)
                    p.paragraph_format.space_after = Pt(3)
            elif "tabela" in b:
                tabela(b["tabela"])
            else:
                g = b.get("grafico") or b.get("grafico_empilhado")
                if g.get("titulo"):
                    paragrafo(g["titulo"], "Texto Agentti").runs[0].bold = True
                doc.add_picture(graficos[id(b)], width=largura)
                if g.get("tabela", True):
                    tabela({**(tabela_do_grafico(g) if "grafico" in b else tabela_do_empilhado(g)),
                            "legenda": g.get("legenda")})
                elif g.get("legenda"):
                    paragrafo(g["legenda"], "Legenda Agentti")

    if c.get("sobre_os_dados", "padrao") is not False:
        paragrafo("Sobre os dados", "Seção Agentti")
        paragrafo(texto_sobre(c), "Texto Agentti")
    doc.core_properties.title = c["titulo"]
    doc.core_properties.author = "Agentti"
    doc.save(caminho)


def texto_sobre(c):
    v = c.get("sobre_os_dados", "padrao")
    if isinstance(v, str) and v not in ("padrao", "sem_contatos"):
        return v
    return SOBRE.format(data=c["data"], contatos="" if v == "sem_contatos" else SOBRE_CONTATOS)


# ── PDF ──────────────────────────────────────────────────────────────────────

def _fontes_pdf():
    """Uma TTF com acentos e símbolos (≥, ×). Sem nenhuma, Helvetica e os símbolos trocados por texto."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pares = [
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
        ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
         "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
         "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf"),
        ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
         "/System/Library/Fonts/Supplemental/Arial Italic.ttf"),
    ]
    try:
        import matplotlib
        d = os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf")
        pares.append((os.path.join(d, "DejaVuSans.ttf"), os.path.join(d, "DejaVuSans-Bold.ttf"),
                      os.path.join(d, "DejaVuSans-Oblique.ttf")))
    except Exception:
        pass
    for reg, neg, ita in pares:
        if all(os.path.exists(p) for p in (reg, neg, ita)):
            pdfmetrics.registerFont(TTFont("AgReg", reg))
            pdfmetrics.registerFont(TTFont("AgNeg", neg))
            pdfmetrics.registerFont(TTFont("AgIta", ita))
            from reportlab.pdfbase.pdfmetrics import registerFontFamily
            registerFontFamily("AgReg", normal="AgReg", bold="AgNeg", italic="AgIta", boldItalic="AgNeg")
            return "AgReg", "AgNeg", "AgIta", False
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", True


def gerar_pdf(c, graficos, caminho):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    reg, neg, ita, simples = _fontes_pdf()
    h = lambda x: colors.HexColor("#" + x)

    def rico(texto):
        t = str(texto)
        if simples:
            t = t.replace("≥", ">=").replace("≤", "<=").replace("×", "x").replace("→", "->")
        return "".join(f"<b>{escape(p)}</b>" if b else escape(p) for p, b in partes(t))

    E = {
        "titulo": ParagraphStyle("t", fontName=neg, fontSize=20, leading=24, textColor=h(LARANJA), spaceAfter=2),
        "sub": ParagraphStyle("s", fontName=reg, fontSize=9, leading=12, textColor=h(SECUNDARIO), spaceAfter=12),
        "secao": ParagraphStyle("se", fontName=neg, fontSize=14, leading=18, textColor=h(LARANJA), spaceBefore=12, spaceAfter=6),
        "texto": ParagraphStyle("x", fontName=reg, fontSize=10.5, leading=14.5, textColor=h(TEXTO), spaceAfter=6),
        "item": ParagraphStyle("i", fontName=reg, fontSize=10.5, leading=14, textColor=h(TEXTO), leftIndent=12, spaceAfter=3),
        "legenda": ParagraphStyle("l", fontName=ita, fontSize=8.5, leading=11, textColor=h(SECUNDARIO), spaceBefore=2, spaceAfter=10),
        "cel": ParagraphStyle("c", fontName=reg, fontSize=9, leading=11.5, textColor=h(TEXTO)),
        "cel_d": ParagraphStyle("cd", fontName=reg, fontSize=9, leading=11.5, textColor=h(TEXTO), alignment=TA_RIGHT),
        "cab": ParagraphStyle("cb", fontName=neg, fontSize=9, leading=11.5, textColor=colors.white),
        "cab_d": ParagraphStyle("cbd", fontName=neg, fontSize=9, leading=11.5, textColor=colors.white, alignment=TA_RIGHT),
        "num": ParagraphStyle("n", fontName=neg, fontSize=18, leading=22, textColor=h(LARANJA), alignment=TA_CENTER),
        "rot": ParagraphStyle("r", fontName=reg, fontSize=8, leading=10, textColor=h(SECUNDARIO), alignment=TA_CENTER),
    }
    largura = A4[0] - 4 * cm

    def tabela(t):
        cols = t["colunas"]
        num = _colunas_numericas(t)
        dados = [[Paragraph(rico(x), E["cab_d"] if j in num else E["cab"]) for j, x in enumerate(cols)]]
        for ln in t["linhas"]:
            dados.append([Paragraph(rico(v), E["cel_d"] if j in num else E["cel"]) for j, v in enumerate(ln)])
        tb = Table(dados, colWidths=[largura / len(cols)] * len(cols), repeatRows=1)
        estilo = [("BACKGROUND", (0, 0), (-1, 0), h(GRAFITE)), ("LINEBELOW", (0, 0), (-1, -1), 0.5, h(LINHA)),
                  ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 4),
                  ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
        estilo += [("BACKGROUND", (0, i), (-1, i), h(FUNDO)) for i in range(2, len(dados), 2)]
        tb.setStyle(TableStyle(estilo))
        out = [tb]
        out.append(Paragraph(rico(t["legenda"]), E["legenda"]) if t.get("legenda") else Spacer(1, 10))
        return out

    fluxo = [Paragraph(rico(c["titulo"]), E["titulo"])]
    if c.get("subtitulo"):
        fluxo.append(Paragraph(rico(c["subtitulo"]), E["sub"]))
    if c.get("destaques"):
        ds = c["destaques"]
        tb = Table([[Paragraph(rico(d["valor"]), E["num"]) for d in ds], [Paragraph(rico(d["rotulo"]), E["rot"]) for d in ds]],
                   colWidths=[largura / len(ds)] * len(ds))
        tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), h(FUNDO)), ("LINEAFTER", (0, 0), (-2, -1), 2, colors.white),
                                ("TOPPADDING", (0, 0), (-1, 0), 8), ("BOTTOMPADDING", (0, -1), (-1, -1), 8)]))
        fluxo += [tb, Spacer(1, 12)]
    for n, s in enumerate(c["secoes"], start=1):
        titulo = s["titulo"] if c.get("numerar") is False else f"{n}. {s['titulo']}"
        corpo = []
        for b in s.get("blocos") or []:
            if "texto" in b:
                corpo.append(Paragraph(rico(b["texto"]), E["texto"]))
            elif "lista" in b:
                corpo += [Paragraph("• " + rico(i), E["item"]) for i in b["lista"]]
            elif "tabela" in b:
                corpo += tabela(b["tabela"])
            else:
                g = b.get("grafico") or b.get("grafico_empilhado")
                img = Image(graficos[id(b)])
                img.drawHeight = largura * img.imageHeight / img.imageWidth
                img.drawWidth = largura
                pedaco = [Paragraph("<b>" + rico(g["titulo"]) + "</b>", E["texto"])] if g.get("titulo") else []
                pedaco.append(img)
                corpo.append(KeepTogether(pedaco))
                if g.get("tabela", True):
                    corpo += tabela({**(tabela_do_grafico(g) if "grafico" in b else tabela_do_empilhado(g)),
                                     "legenda": g.get("legenda")})
                elif g.get("legenda"):
                    corpo.append(Paragraph(rico(g["legenda"]), E["legenda"]))
        fluxo.append(KeepTogether([Paragraph(rico(titulo), E["secao"])] + corpo[:1]))
        fluxo += corpo[1:]
    if c.get("sobre_os_dados", "padrao") is not False:
        fluxo += [Paragraph("Sobre os dados", E["secao"]), Paragraph(rico(texto_sobre(c)), E["texto"])]

    curto = c.get("titulo_curto") or c["titulo"]

    def moldura(canv, _doc):
        canv.saveState()
        topo = A4[1] - 1.2 * cm
        if os.path.exists(LOGO):
            canv.drawImage(LOGO, 2 * cm, topo - 1.0 * cm, width=2.2 * cm, height=1.42 * cm, mask="auto",
                           preserveAspectRatio=True, anchor="sw")
        canv.setFont(reg, 8)
        canv.setFillColor(h(SECUNDARIO))
        canv.drawRightString(A4[0] - 2 * cm, topo - 0.6 * cm, curto)
        rod = f"{RODAPE} · gerado em {c['data']}"
        canv.drawString(2 * cm, 1.2 * cm, rod if not simples else rod.replace("·", "-"))
        canv.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Página {canv.getPageNumber()}")
        canv.restoreState()

    SimpleDocTemplate(caminho, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2.8 * cm,
                      bottomMargin=2 * cm, title=c["titulo"], author="Agentti").build(fluxo, onFirstPage=moldura,
                                                                                      onLaterPages=moldura)


# ── Principal ────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("conteudo")
    ap.add_argument("--saida", default=".")
    ap.add_argument("--previa", action="store_true", help="também um PNG pequeno da 1ª página, para conferir")
    a = ap.parse_args()

    faltam = []
    for mod, pacote in (("docx", "python-docx"), ("matplotlib", "matplotlib"), ("reportlab", "reportlab")):
        try:
            __import__(mod)
        except ImportError:
            faltam.append(pacote)
    if faltam:
        print(json.dumps({"erro": "faltam pacotes", "instalar": "pip install " + " ".join(faltam)}, ensure_ascii=False))
        sys.exit(3)

    try:
        # utf-8-sig: aceita o BOM que editores do Windows colocam no começo do arquivo.
        with open(a.conteudo, encoding="utf-8-sig") as f:
            c = json.load(f)
    except FileNotFoundError:
        print(json.dumps({"erro": f"arquivo não encontrado: {a.conteudo}"}, ensure_ascii=False))
        sys.exit(2)
    except json.JSONDecodeError as e:
        print(json.dumps({"erro": "JSON inválido", "detalhes": [f"linha {e.lineno}, coluna {e.colno}: {e.msg}"]},
                         ensure_ascii=False))
        sys.exit(2)
    erros = validar(c)
    if erros:
        print(json.dumps({"erro": "conteúdo inválido", "detalhes": erros}, ensure_ascii=False))
        sys.exit(2)
    c.setdefault("data", date.today().strftime("%d/%m/%Y"))
    nome = re.sub(r'[<>:"/\\|?*]', "-", c.get("arquivo") or f"Agentti - {c['titulo']}").strip()
    os.makedirs(a.saida, exist_ok=True)

    tmp = tempfile.mkdtemp(prefix="agentti_")
    graficos = {}
    for s in c["secoes"]:
        for b in s.get("blocos") or []:
            if "grafico" in b or "grafico_empilhado" in b:
                p = os.path.join(tmp, f"g{len(graficos)}.png")
                (grafico_barras if "grafico" in b else grafico_empilhado)(b.get("grafico") or b["grafico_empilhado"], p)
                graficos[id(b)] = p

    docx_p = os.path.join(a.saida, nome + ".docx")
    pdf_p = os.path.join(a.saida, nome + ".pdf")
    gerar_docx(c, graficos, docx_p)
    gerar_pdf(c, graficos, pdf_p)
    saida = {"docx": os.path.abspath(docx_p), "pdf": os.path.abspath(pdf_p), "graficos": len(graficos)}
    if a.previa:
        try:
            try:
                import pymupdf as fitz
            except ImportError:
                import fitz
            doc = fitz.open(pdf_p)
            saida["paginas"] = doc.page_count
            pv = os.path.join(a.saida, nome + " - previa.png")
            doc[0].get_pixmap(dpi=45).save(pv)
            saida["previa"] = os.path.abspath(pv)
        except Exception as e:
            saida["previa"] = f"sem prévia ({type(e).__name__}); os arquivos foram gerados"
    print(json.dumps(saida, ensure_ascii=False))


if __name__ == "__main__":
    main()
